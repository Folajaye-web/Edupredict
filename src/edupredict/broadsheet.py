"""Phase 3.2 intelligent broadsheet services.

All views are derived from placements, subject enrolments, results and attendance.
No collated/broadsheet table is persisted.  The module deliberately distinguishes:
- recorded result: numeric score/grade
- enrolled but result missing: em dash (—)
- not enrolled: N/E

Aggregates only use recorded scores; missing results never become zero/F9.
"""
from __future__ import annotations

from io import BytesIO
from typing import Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .analytics import analytics_dataset
from .db import AcademicSession
from .grading import grade_from_score
from .analytics import DEFAULT_CREDIT_THRESHOLD, DEFAULT_PASS_THRESHOLD


IDENTITY_COLUMNS = ["S/N", "Student ID", "Student Name", "Class", "Arm / Stream"]
SUMMARY_COLUMNS = [
    "Average", "Performance Band", "Rank", "Ranking Status", "Subjects Enrolled", "Subjects Recorded",
    "Missing Results", "Completion %", "Below Credit", "Credit Pass %", "Attendance %",
]


def _period(s: Session, school_id: int, session_name: str, term: str):
    return s.query(AcademicSession).filter_by(
        school_id=school_id, session_name=session_name, term=term
    ).first()


def _empty_frame():
    return pd.DataFrame(columns=IDENTITY_COLUMNS + SUMMARY_COLUMNS)


def build_broadsheet(
    s: Session,
    school_id: int,
    session_name: str,
    term: str,
    class_level: Optional[str] = None,
    arm: Optional[str] = None,
    subject_code: Optional[str] = None,
    credit_threshold: float = DEFAULT_CREDIT_THRESHOLD,
    pass_threshold: float = DEFAULT_PASS_THRESHOLD,
) -> pd.DataFrame:
    """Build the canonical numeric broadsheet from current transactional data."""
    sess = _period(s, school_id, session_name, term)
    if not sess:
        return _empty_frame()

    df = analytics_dataset(
        s, school_id, sess.id, class_level, arm,
        credit_threshold=credit_threshold, pass_threshold=pass_threshold,
    )
    if df.empty:
        return _empty_frame()
    if subject_code:
        df = df[df["subject_code"] == subject_code].copy()
    if df.empty:
        return _empty_frame()

    # Subject order is deterministic and curriculum-friendly: catalogue name, then code.
    subject_order = (
        df[["subject_code", "subject"]]
        .drop_duplicates()
        .sort_values(["subject", "subject_code"], kind="stable")
    )
    subject_names = subject_order["subject"].tolist()

    # One row per placed student. Each student gets the union of subjects offered to the
    # selected cohort; cells distinguish N/E from enrolled-but-missing later in display mode.
    students = (
        df[["student_id", "student_code", "student_name", "class_level", "arm"]]
        .drop_duplicates("student_id")
        .sort_values(["class_level", "arm", "student_code"], kind="stable")
    )
    enrolled_counts = df.groupby("student_id")["subject_id"].nunique()
    recorded_counts = df.groupby("student_id")["result_recorded"].sum()
    scores = df.pivot_table(
        index="student_id", columns="subject", values="total", aggfunc="first"
    )
    scores = scores.reindex(index=students.student_id, columns=subject_names)

    out = pd.DataFrame({
        "student_id": students.student_id.to_numpy(),
        "Student ID": students.student_code.to_numpy(),
        "Student Name": students.student_name.to_numpy(),
        "Class": students.class_level.to_numpy(),
        "Arm / Stream": students.arm.replace("", "General").to_numpy(),
    }).set_index("student_id")
    for subject in subject_names:
        out[subject] = scores[subject]

    numeric = out[subject_names].apply(pd.to_numeric, errors="coerce")
    out["Subjects Enrolled"] = enrolled_counts.reindex(out.index).fillna(0).astype(int)
    out["Subjects Recorded"] = recorded_counts.reindex(out.index).fillna(0).astype(int)
    out["Missing Results"] = out["Subjects Enrolled"] - out["Subjects Recorded"]
    out["Completion %"] = (
        out["Subjects Recorded"] / out["Subjects Enrolled"].replace(0, np.nan) * 100
    ).round(1)
    out["Average"] = numeric.mean(axis=1).round(1)
    out["Performance Band"] = out["Average"].map(
        lambda x: grade_from_score(x)[0] if pd.notna(x) else "No result"
    )
    out["Below Credit"] = numeric.lt(credit_threshold).sum(axis=1).astype(int)
    recorded = numeric.notna().sum(axis=1)
    out["Credit Pass %"] = (
        numeric.ge(credit_threshold).sum(axis=1) / recorded.replace(0, np.nan) * 100
    ).round(1)
    att = df.groupby("student_id")["attendance_pct"].first()
    out["Attendance %"] = att.reindex(out.index).round(1)

    # Descriptive rank: only students with at least one recorded score can rank.
    out["Rank"] = out["Average"].rank(method="min", ascending=False).astype("Int64")
    out["Ranking Status"] = np.where(out["Average"].notna(), np.where(out["Completion %"].eq(100), "Ranked — complete", "Ranked — incomplete"), "Not ranked — no result")
    out["S/N"] = np.arange(1, len(out) + 1)
    out = out.reset_index(drop=True)
    return out[["S/N", "Student ID", "Student Name", "Class", "Arm / Stream"] + subject_names + SUMMARY_COLUMNS]


def build_broadsheet_display(
    s: Session,
    school_id: int,
    session_name: str,
    term: str,
    class_level: Optional[str] = None,
    arm: Optional[str] = None,
    subject_code: Optional[str] = None,
    credit_threshold: float = DEFAULT_CREDIT_THRESHOLD,
    pass_threshold: float = DEFAULT_PASS_THRESHOLD,
) -> pd.DataFrame:
    """Return the teacher-facing matrix with N/E and — state markers."""
    numeric = build_broadsheet(
        s, school_id, session_name, term, class_level, arm, subject_code,
        credit_threshold, pass_threshold,
    )
    if numeric.empty:
        return numeric
    sess = _period(s, school_id, session_name, term)
    raw = analytics_dataset(
        s, school_id, sess.id, class_level, arm,
        credit_threshold=credit_threshold, pass_threshold=pass_threshold,
    )
    if subject_code:
        raw = raw[raw["subject_code"] == subject_code]
    enrollment = {(int(r.student_id), r.subject) for r in raw.itertuples()}
    recorded = {
        (int(r.student_id), r.subject)
        for r in raw[raw.result_recorded].itertuples()
    }
    # Recover stable student IDs in display order.
    id_by_code = dict(zip(numeric["Student ID"], raw.drop_duplicates("student_id").sort_values(["class_level", "arm", "student_code"])["student_id"]))
    display = numeric.copy()
    subject_cols = [c for c in numeric.columns if c not in IDENTITY_COLUMNS + SUMMARY_COLUMNS]
    display[subject_cols] = display[subject_cols].astype(object)
    for i, row in display.iterrows():
        sid = id_by_code.get(row["Student ID"])
        for subject in subject_cols:
            key = (int(sid), subject)
            if key not in enrollment:
                display.at[i, subject] = "N/E"
            elif key not in recorded:
                display.at[i, subject] = "—"
            else:
                value = pd.to_numeric(display.at[i, subject], errors="coerce")
                display.at[i, subject] = f"{value:.1f}" if pd.notna(value) else "—"
    for col in ["Average", "Completion %", "Credit Pass %", "Attendance %"]:
        display[col] = display[col].map(lambda x: f"{x:.1f}" if pd.notna(x) else "—")
    return display


def broadsheet_missing_results(
    s: Session, school_id: int, session_name: str, term: str,
    class_level: Optional[str] = None, arm: Optional[str] = None,
) -> pd.DataFrame:
    """List every enrolled student-subject pair with no current result."""
    sess = _period(s, school_id, session_name, term)
    if not sess:
        return pd.DataFrame()
    df = analytics_dataset(s, school_id, sess.id, class_level, arm)
    if df.empty:
        return pd.DataFrame(columns=["Student ID", "Student Name", "Class", "Arm / Stream", "Subject Code", "Subject"])
    missing = df[~df.result_recorded].copy()
    if missing.empty:
        return pd.DataFrame(columns=["Student ID", "Student Name", "Class", "Arm / Stream", "Subject Code", "Subject"])
    return (
        missing[["student_code", "student_name", "class_level", "arm", "subject_code", "subject"]]
        .rename(columns={
            "student_code": "Student ID", "student_name": "Student Name",
            "class_level": "Class", "arm": "Arm / Stream",
            "subject_code": "Subject Code", "subject": "Subject",
        })
        .sort_values(["Class", "Arm / Stream", "Student ID", "Subject"])
        .reset_index(drop=True)
    )


def broadsheet_subject_summary(
    s: Session, school_id: int, session_name: str, term: str,
    class_level: Optional[str] = None, arm: Optional[str] = None,
    credit_threshold: float = DEFAULT_CREDIT_THRESHOLD,
    pass_threshold: float = DEFAULT_PASS_THRESHOLD,
) -> pd.DataFrame:
    """Teacher-oriented subject summary including enrollment and completion."""
    sess = _period(s, school_id, session_name, term)
    if not sess:
        return pd.DataFrame()
    from .analytics import subject_analytics
    return subject_analytics(
        s, school_id, sess.id, class_level, arm,
        credit_threshold=credit_threshold, pass_threshold=pass_threshold,
    )


def broadsheet_kpis(df: pd.DataFrame, credit_threshold: float = DEFAULT_CREDIT_THRESHOLD) -> dict:
    """Calculate class-level broadsheet KPIs from the canonical frame."""
    if df.empty:
        return {
            "students": 0, "subjects": 0, "enrolled_records": 0, "recorded_results": 0,
            "missing_results": 0, "completion_rate": 0.0, "average": None,
            "credit_rate": None, "pass_rate": None, "attendance": None,
            "students_below_credit": 0, "students_with_complete_results": 0,
        }
    subject_cols = [c for c in df.columns if c not in IDENTITY_COLUMNS + SUMMARY_COLUMNS + ["student_id"]]
    scores = df[subject_cols].apply(pd.to_numeric, errors="coerce")
    observed = pd.Series(scores.to_numpy(dtype=float).ravel()).dropna()
    avg = pd.to_numeric(df["Average"], errors="coerce")
    complete = pd.to_numeric(df["Completion %"], errors="coerce").eq(100)
    return {
        "students": int(len(df)),
        "subjects": int(len(subject_cols)),
        "enrolled_records": int(pd.to_numeric(df["Subjects Enrolled"], errors="coerce").sum()),
        "recorded_results": int(pd.to_numeric(df["Subjects Recorded"], errors="coerce").sum()),
        "missing_results": int(pd.to_numeric(df["Missing Results"], errors="coerce").sum()),
        "completion_rate": round(float(observed.size / max(1, pd.to_numeric(df["Subjects Enrolled"], errors="coerce").sum()) * 100), 1),
        "average": round(float(observed.mean()), 2) if not observed.empty else None,
        "credit_rate": round(float(observed.ge(credit_threshold).mean() * 100), 1) if not observed.empty else None,
        "pass_rate": round(float(observed.ge(DEFAULT_PASS_THRESHOLD).mean() * 100), 1) if not observed.empty else None,
        "attendance": round(float(pd.to_numeric(df["Attendance %"], errors="coerce").dropna().mean()), 1) if pd.to_numeric(df["Attendance %"], errors="coerce").notna().any() else None,
        "students_below_credit": int(avg.lt(credit_threshold).sum()),
        "students_with_complete_results": int(complete.sum()),
    }


def export_broadsheet_xlsx(
    numeric_df: pd.DataFrame,
    display_df: Optional[pd.DataFrame] = None,
    subject_summary: Optional[pd.DataFrame] = None,
    missing_results: Optional[pd.DataFrame] = None,
) -> bytes:
    """Create a clean multi-sheet Excel workbook for school reporting."""
    if display_df is None:
        display_df = numeric_df
    wb = Workbook()
    ws = wb.active
    ws.title = "Broadsheet"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    def write_df(sheet, frame):
        sheet.append(list(frame.columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for values in frame.itertuples(index=False, name=None):
            sheet.append(list(values))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col_cells in sheet.columns:
            max_len = min(35, max(len(str(c.value or "")) for c in col_cells) + 2)
            sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = max(10, max_len)

    write_df(ws, display_df)
    if subject_summary is not None and not subject_summary.empty:
        write_df(wb.create_sheet("Subject Summary"), subject_summary)
    if missing_results is not None and not missing_results.empty:
        write_df(wb.create_sheet("Missing Results"), missing_results)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
